import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = ["No.", "Name", "Pet", "Type", "OPD", "Description", "Price", "Note"]
KEYS    = ["no",  "name", "pet", "type", "opd", "description", "price", "note"]

HEADER_FONT  = Font(name="Arial", bold=True, size=11, color="FFFFFF")
HEADER_FILL  = PatternFill("solid", fgColor="2E75B6")
ROW_FILL_ODD  = PatternFill("solid", fgColor="FFFFFF")
ROW_FILL_EVEN = PatternFill("solid", fgColor="EEF3FA")
DATA_FONT     = Font(name="Arial", size=10)
PRICE_FONT    = Font(name="Arial", size=10)

_thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

COL_WIDTHS = {
    1: 6,   # No.
    2: 18,  # Name
    3: 14,  # Pet
    4: 10,  # Type
    5: 10,  # OPD
    6: 45,  # Description
    7: 10,  # Price
    8: 20,  # Note
}


def build_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Patient Records"

    # Header row
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    # Data rows — renumber sequentially
    for row_idx, record in enumerate(rows, start=1):
        excel_row = row_idx + 1
        fill = ROW_FILL_ODD if row_idx % 2 == 1 else ROW_FILL_EVEN

        for col_idx, key in enumerate(KEYS, start=1):
            value = record.get(key, "")
            if key == "no":
                value = row_idx  # renumber sequentially
            elif key == "price":
                try:
                    value = float(value) if value not in ("", None) else ""
                except (ValueError, TypeError):
                    pass  # keep as string if not numeric

            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.font = PRICE_FONT if key == "price" else DATA_FONT
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
                horizontal="right" if key in ("no", "price") else "left",
            )

        if record.get("description", "") and len(record["description"]) > 60:
            ws.row_dimensions[excel_row].height = 36
        else:
            ws.row_dimensions[excel_row].height = 18

    # Column widths
    for col_idx, width in COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Price column: number format
    price_col = get_column_letter(KEYS.index("price") + 1)
    for row in ws.iter_rows(min_row=2, min_col=KEYS.index("price") + 1,
                             max_col=KEYS.index("price") + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"

    # Freeze header row
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
